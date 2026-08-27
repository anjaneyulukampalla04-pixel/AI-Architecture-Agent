import os
import shutil
import asyncio

from pathlib import Path
from typing import List, Optional, Dict, Any

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    WebSocket,
    WebSocketDisconnect,
    File,
    UploadFile,
)

from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from app.database import get_db
from app.models import (
    Project,
    Deployment,
    User,
    Document,
    ArchitectureVersion,
    Review,
    ReviewFinding,
    Requirement,
)

from app.schemas import (
    ProjectCreate,
    ProjectUpdate,
    ProjectResponse,
    ProjectsListResponse,
    DeploymentCreate,
    FindingUpdate,
)

from app.middleware.auth_middleware import get_current_user
from app.middleware.rbac_middleware import require_role
from app.websocket import manager
from app.config import settings


router = APIRouter(
    prefix="/projects",
    tags=["projects"]
)


# ============================================================
# UPLOAD DIRECTORY
# ============================================================

os.makedirs(
    settings.UPLOAD_DIR,
    exist_ok=True
)


# ============================================================
# DOCUMENT TEXT EXTRACTION
# ============================================================

async def extract_text_from_file(file_path: str) -> str:
    """
    Extract readable text from uploaded documents.

    Supported:
    - PDF
    - DOCX
    - TXT
    - MD
    - JSON
    - CSV
    - Source-code/text files

    IMPORTANT:
    Binary files such as PDFs must never be opened directly
    as UTF-8 text because doing so produces raw PDF internals
    such as /Type /Page, endobj, startxref and %%EOF.
    """

    path = Path(file_path)

    extension = path.suffix.lower()

    try:

        # ====================================================
        # PDF
        # ====================================================

        if extension == ".pdf":

            from pypdf import PdfReader

            reader = PdfReader(file_path)

            pages = []

            for page_number, page in enumerate(
                reader.pages,
                start=1
            ):

                try:

                    text = page.extract_text()

                    if text and text.strip():

                        pages.append(
                            text.strip()
                        )

                except Exception as page_error:

                    print(
                        f"[DOCUMENT] Could not extract "
                        f"page {page_number}: "
                        f"{repr(page_error)}"
                    )

            extracted_text = "\n\n".join(pages)

            if not extracted_text.strip():

                return (
                    "[PDF uploaded successfully, "
                    "but no readable text could be extracted. "
                    "The PDF may contain scanned images rather "
                    "than selectable text.]"
                )

            print(
                f"[DOCUMENT] PDF extracted successfully: "
                f"{path.name}"
            )

            print(
                f"[DOCUMENT] Extracted characters: "
                f"{len(extracted_text)}"
            )

            return extracted_text.strip()


        # ====================================================
        # DOCX
        # ====================================================

        elif extension == ".docx":

            from docx import Document as DocxDocument

            document = DocxDocument(
                file_path
            )

            content = []

            # ------------------------------------------------
            # Paragraphs
            # ------------------------------------------------

            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:

                    content.append(text)


            # ------------------------------------------------
            # Tables
            # ------------------------------------------------

            for table in document.tables:

                for row in table.rows:

                    row_text = []

                    for cell in row.cells:

                        value = cell.text.strip()

                        if value:

                            row_text.append(value)

                    if row_text:

                        content.append(
                            " | ".join(row_text)
                        )


            extracted_text = "\n".join(content)

            if not extracted_text.strip():

                return (
                    "[DOCX uploaded successfully, "
                    "but no readable text could be extracted.]"
                )

            print(
                f"[DOCUMENT] DOCX extracted successfully: "
                f"{path.name}"
            )

            return extracted_text.strip()


        # ====================================================
        # NORMAL TEXT FILES
        # ====================================================

        elif extension in {

            ".txt",
            ".md",
            ".json",
            ".csv",

            ".py",
            ".js",
            ".jsx",

            ".ts",
            ".tsx",

            ".html",
            ".css",

            ".yaml",
            ".yml",

            ".xml",

            ".tf",
            ".tfvars",

            ".sql",

        }:

            with open(
                file_path,
                "r",
                encoding="utf-8",
                errors="replace"
            ) as file:

                text = file.read()

            print(
                f"[DOCUMENT] Text file extracted: "
                f"{path.name}"
            )

            return text.strip()


        # ====================================================
        # UNSUPPORTED FILE
        # ====================================================

        else:

            print(
                f"[DOCUMENT] Unsupported file type: "
                f"{extension}"
            )

            return (
                f"[Document uploaded successfully: "
                f"{path.name}. "
                f"Automatic text extraction is not "
                f"supported for file type "
                f"{extension or 'unknown'}.]"
            )


    except Exception as error:

        print("=" * 70)

        print(
            "[DOCUMENT] TEXT EXTRACTION FAILED"
        )

        print(
            f"[DOCUMENT] File: {path.name}"
        )

        print(
            f"[DOCUMENT] Error: {repr(error)}"
        )

        print("=" * 70)

        return (
            f"[Could not extract readable text "
            f"from {path.name}.]"
        )


# ============================================================
# PROJECTS
# ============================================================

@router.get(
    "",
    response_model=ProjectsListResponse
)
async def get_projects(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(
        select(Project)
        .filter(
            Project.ownerId == current_user.id
        )
    )

    projects = result.scalars().all()

    return {
        "count": len(projects),
        "projects": projects
    }


@router.post(
    "",
    response_model=ProjectResponse,
    status_code=201
)
async def create_project(
    project: ProjectCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    new_project = Project(
        name=project.name,
        description=project.description,
        ownerId=current_user.id
    )

    db.add(new_project)

    await db.commit()

    await db.refresh(new_project)

    return new_project


@router.get(
    "/{id}",
    response_model=ProjectResponse
)
async def get_project(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(
        select(Project)
        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )

    return project


@router.put(
    "/{id}",
    response_model=ProjectResponse
)
async def update_project(
    id: str,
    project_update: ProjectUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(
        select(Project)
        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )

    if project_update.description is not None:

        project.description = (
            project_update.description
        )

    await db.commit()

    await db.refresh(project)

    return project


@router.delete(
    "/{id}",
    status_code=204
)
async def delete_project(
    id: str,
    current_user: User = Depends(
        require_role("admin")
    ),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(
        select(Project)
        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )

    await db.delete(project)

    await db.commit()

    return None


# ============================================================
# DOCUMENT UPLOAD
# ============================================================

@router.post(
    "/{id}/documents",
    status_code=201
)
async def upload_document(
    id: str,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    # --------------------------------------------------------
    # Verify ownership
    # --------------------------------------------------------

    result = await db.execute(
        select(Project)
        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    # --------------------------------------------------------
    # Validate filename
    # --------------------------------------------------------

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="Uploaded file has no filename"
        )


    # --------------------------------------------------------
    # Save file
    # --------------------------------------------------------

    safe_filename = os.path.basename(
        file.filename
    )

    file_path = os.path.join(
        settings.UPLOAD_DIR,
        f"{id}_{safe_filename}"
    )


    try:

        with open(
            file_path,
            "wb"
        ) as buffer:

            shutil.copyfileobj(
                file.file,
                buffer
            )

    except Exception as error:

        print(
            "[DOCUMENT] Failed to save upload:",
            repr(error)
        )

        raise HTTPException(
            status_code=500,
            detail="Failed to save uploaded document"
        )


    # --------------------------------------------------------
    # File size
    # --------------------------------------------------------

    file_size = os.path.getsize(
        file_path
    )


    # --------------------------------------------------------
    # Extract actual readable text
    # --------------------------------------------------------

    extracted_text = (
        await extract_text_from_file(
            file_path
        )
    )


    # --------------------------------------------------------
    # Prevent accidental PDF internals from being stored
    # --------------------------------------------------------

    raw_pdf_markers = [
        "%PDF-",
        "/Type /Page",
        "/MediaBox",
        "startxref",
        "%%EOF",
    ]

    marker_count = sum(
        marker in extracted_text
        for marker in raw_pdf_markers
    )

    if marker_count >= 2:

        print(
            "[DOCUMENT] WARNING: Raw PDF "
            "structure detected. Discarding "
            "extracted text."
        )

        extracted_text = (
            "[PDF uploaded, but readable text "
            "could not be safely extracted.]"
        )


    # --------------------------------------------------------
    # Save document metadata + extracted text
    # --------------------------------------------------------

    new_doc = Document(
        projectId=id,
        name=safe_filename,
        filePath=file_path,
        fileSize=file_size,
        extractedText=extracted_text
    )

    db.add(new_doc)

    await db.commit()

    await db.refresh(new_doc)


    return {

        "id": new_doc.id,

        "name": new_doc.name,

        "fileSize": new_doc.fileSize,

        "message":
            "Document uploaded and parsed successfully!"
    }


# ============================================================
# PROJECT DOCUMENTS
# ============================================================

@router.get(
    "/{id}/documents"
)
async def get_project_documents(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(
        select(Project)
        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    docs_result = await db.execute(

        select(Document)

        .filter(
            Document.projectId == id
        )

        .order_by(
            Document.createdAt.desc()
        )
    )

    docs = docs_result.scalars().all()


    return [

        {
            "id": document.id,

            "name": document.name,

            "fileSize": document.fileSize,

            "createdAt": document.createdAt
        }

        for document in docs
    ]


# ============================================================
# ARCHITECTURE
# ============================================================

@router.get(
    "/{id}/architecture"
)
async def get_logical_architecture(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    latest_version = (
        result.scalars().first()
    )

    if (
        not latest_version
        or not latest_version.architectureData
    ):

        raise HTTPException(
            status_code=404,
            detail=(
                "No architecture design "
                "generated yet."
            )
        )

    return (
        latest_version
        .architectureData
        .get("components", [])
    )


# ============================================================
# DATABASE SCHEMA
# ============================================================

@router.get(
    "/{id}/architecture/database"
)
async def get_database_schema(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    latest_version = (
        result.scalars().first()
    )

    if (
        not latest_version
        or not latest_version.architectureData
    ):

        raise HTTPException(
            status_code=404,
            detail=(
                "No database schema generated yet."
            )
        )

    return (
        latest_version
        .architectureData
        .get(
            "database_schema",
            []
        )
    )


# ============================================================
# API SPECIFICATIONS
# ============================================================

@router.get(
    "/{id}/architecture/apis"
)
async def get_api_specifications(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    latest_version = (
        result.scalars().first()
    )

    if (
        not latest_version
        or not latest_version.architectureData
    ):

        raise HTTPException(
            status_code=404,
            detail=(
                "No API specifications generated yet."
            )
        )

    return (
        latest_version
        .architectureData
        .get("apis", [])
    )


# ============================================================
# CLOUD MAPPING
# ============================================================

@router.get(
    "/{id}/cloud-mapping"
)
async def get_cloud_mappings(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    latest_version = (
        result.scalars().first()
    )

    if (
        not latest_version
        or not latest_version.architectureData
    ):

        raise HTTPException(
            status_code=404,
            detail=(
                "No cloud service mapping generated yet."
            )
        )


    return {

        "mappings":
            latest_version
            .architectureData
            .get(
                "cloud_mappings",
                []
            ),

        "costs":
            latest_version
            .architectureData
            .get(
                "costs",
                {}
            )
    }


# ============================================================
# TERRAFORM
# ============================================================

@router.get(
    "/{id}/terraform"
)
async def get_terraform_code(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    latest_version = (
        result.scalars().first()
    )

    if (
        not latest_version
        or not latest_version.terraformCode
    ):

        raise HTTPException(
            status_code=404,
            detail=(
                "No Terraform code compiled yet."
            )
        )


    return {

        "version":
            latest_version.versionNumber,

        "terraformCode":
            latest_version.terraformCode
    }


# ============================================================
# REVIEWS
# ============================================================

@router.get(
    "/{id}/reviews"
)
async def get_reviews_findings(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    review_res = await db.execute(

        select(Review)

        .filter(
            Review.projectId == id
        )

        .order_by(
            Review.createdAt.desc()
        )
    )

    latest_review = (
        review_res.scalars().first()
    )

    if not latest_review:

        raise HTTPException(
            status_code=404,
            detail=(
                "No architecture audit "
                "grading performed yet."
            )
        )


    findings_res = await db.execute(

        select(ReviewFinding)

        .filter(
            ReviewFinding.reviewId
            == latest_review.id
        )
    )

    findings = (
        findings_res.scalars().all()
    )


    return {

        "score":
            latest_review.score,

        "status":
            latest_review.status,

        "findings":
            findings
    }


# ============================================================
# VERSIONS
# ============================================================

@router.get(
    "/{id}/versions"
)
async def get_all_versions(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    versions = result.scalars().all()


    return [

        {
            "id": version.id,

            "versionNumber":
                version.versionNumber,

            "createdAt":
                version.createdAt
        }

        for version in versions
    ]


# ============================================================
# PROJECT MESSAGES
# ============================================================

@router.get(
    "/{id}/messages"
)
async def get_project_messages(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    from app.models import (
        Message,
        Conversation
    )


    result = await db.execute(

        select(Project)

        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    conv_result = await db.execute(

        select(Conversation)

        .filter(
            Conversation.projectId == id
        )

        .order_by(
            Conversation.createdAt.asc()
        )
    )

    conversation = (
        conv_result.scalars().first()
    )

    if not conversation:

        return []


    msg_result = await db.execute(

        select(Message)

        .filter(
            Message.conversationId
            == conversation.id
        )

        .order_by(
            Message.createdAt.asc()
        )
    )

    messages = (
        msg_result.scalars().all()
    )


    return [

        {
            "sender": message.sender,

            "content": message.content,

            "createdAt":
                message.createdAt
        }

        for message in messages
    ]


# ============================================================
# LIVE SPECIFICATION EDITING
# ============================================================

class DocumentationUpdate(BaseModel):

    markdown: str


class ArchitectureUpdatePayload(BaseModel):

    architectureData: Dict[str, Any]

    terraformCode: Optional[str] = None


# ============================================================
# DOCUMENTATION
# ============================================================

@router.get(
    "/{id}/documentation"
)
async def get_project_documentation(
    id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(Project)

        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    docs_path = os.path.abspath(

        os.path.join(

            os.path.dirname(__file__),

            "../../../docs/architecture_docs.md"
        )
    )


    if not os.path.exists(
        docs_path
    ):

        return {

            "markdown":
                "# Executive Architecture Specifications\n\n"
                "No documentation compiled yet. "
                "Start the pipeline to generate it!"
        }


    try:

        with open(
            docs_path,
            "r",
            encoding="utf-8"
        ) as file:

            return {
                "markdown":
                    file.read()
            }

    except Exception as error:

        raise HTTPException(

            status_code=500,

            detail=(
                "Failed to read documentation: "
                f"{str(error)}"
            )
        )


@router.post(
    "/{id}/documentation"
)
async def update_project_documentation(
    id: str,
    payload: DocumentationUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(Project)

        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    docs_path = os.path.abspath(

        os.path.join(

            os.path.dirname(__file__),

            "../../../docs/architecture_docs.md"
        )
    )


    os.makedirs(
        os.path.dirname(docs_path),
        exist_ok=True
    )


    try:

        with open(
            docs_path,
            "w",
            encoding="utf-8"
        ) as file:

            file.write(
                payload.markdown
            )


        return {

            "status": "success",

            "message":
                "Documentation saved successfully!"
        }


    except Exception as error:

        raise HTTPException(

            status_code=500,

            detail=(
                "Failed to write documentation: "
                f"{str(error)}"
            )
        )


# ============================================================
# UPDATE ARCHITECTURE
# ============================================================

@router.put(
    "/{id}/architecture"
)
async def update_project_architecture(
    id: str,
    payload: ArchitectureUpdatePayload,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(Project)

        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    ver_result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id
        )

        .order_by(
            ArchitectureVersion.versionNumber.desc()
        )
    )

    latest_ver = (
        ver_result.scalars().first()
    )


    if not latest_ver:

        raise HTTPException(
            status_code=404,
            detail="No architecture versions found"
        )


    latest_ver.architectureData = (
        payload.architectureData
    )


    if payload.terraformCode is not None:

        latest_ver.terraformCode = (
            payload.terraformCode
        )


    await db.commit()


    return {

        "status": "success",

        "message":
            "Architecture updated successfully!"
    }


# ============================================================
# DEPLOYMENT
# ============================================================

@router.post(
    "/{id}/deploy"
)
async def deploy_project(
    id: str,
    deployment_data: DeploymentCreate,
    current_user: User = Depends(
        require_role("admin")
    ),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(Project)

        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    new_deployment = Deployment(

        projectId=id,

        environment=
            deployment_data.environment,

        status="pending"
    )


    db.add(new_deployment)

    await db.commit()

    await db.refresh(new_deployment)


    asyncio.create_task(

        mock_deployment_task(
            id,
            new_deployment.id
        )
    )


    return {

        "status": "success",

        "deploymentId":
            new_deployment.id
    }


async def mock_deployment_task(
    project_id: str,
    deployment_id: str
):

    await asyncio.sleep(1)


    await manager.broadcast(

        project_id,

        {
            "type":
                "deployment_update",

            "data": {

                "deploymentId":
                    deployment_id,

                "status":
                    "in_progress",

                "log":
                    "Provisioning resources..."
            }
        }
    )


    await asyncio.sleep(2)


    await manager.broadcast(

        project_id,

        {
            "type":
                "deployment_update",

            "data": {

                "deploymentId":
                    deployment_id,

                "status":
                    "completed",

                "log":
                    "Deployment successful."
            }
        }
    )


# ============================================================
# FINDINGS
# ============================================================

@router.put(
    "/{id}/findings/{findingId}"
)
async def update_finding(
    id: str,
    findingId: str,
    finding_update: FindingUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    return {

        "id": findingId,

        "reviewId":
            "mock_review_id",

        "ruleId":
            "SEC-001",

        "severity":
            "high",

        "description":
            "Mock finding",

        "status":
            finding_update.status,

        "recommendation":
            "Mock recommendation",

        "createdAt":
            "2023-01-01T00:00:00Z",

        "updatedAt":
            "2023-01-01T00:00:00Z"
    }


# ============================================================
# WEBSOCKET
# ============================================================

@router.websocket(
    "/{id}/ws"
)
async def websocket_endpoint(
    websocket: WebSocket,
    id: str
):

    await manager.connect(
        websocket,
        id
    )

    try:

        while True:

            await websocket.receive_text()

    except WebSocketDisconnect:

        manager.disconnect(
            websocket,
            id
        )


# ============================================================
# VERSION COMPARISON
# ============================================================

@router.get(
    "/{id}/compare"
)
async def compare_versions(
    id: str,
    v1: int,
    v2: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):

    result = await db.execute(

        select(Project)

        .filter(
            Project.id == id,
            Project.ownerId == current_user.id
        )
    )

    project = result.scalars().first()

    if not project:

        raise HTTPException(
            status_code=404,
            detail="Project not found"
        )


    v1_result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id,
            ArchitectureVersion.versionNumber == v1
        )
    )

    v1_version = (
        v1_result.scalars().first()
    )


    if not v1_version:

        raise HTTPException(
            status_code=404,
            detail=f"Version {v1} not found"
        )


    v2_result = await db.execute(

        select(ArchitectureVersion)

        .filter(
            ArchitectureVersion.projectId == id,
            ArchitectureVersion.versionNumber == v2
        )
    )

    v2_version = (
        v2_result.scalars().first()
    )


    if not v2_version:

        raise HTTPException(
            status_code=404,
            detail=f"Version {v2} not found"
        )


    v1_data = (
        v1_version.architectureData
        or {}
    )

    v2_data = (
        v2_version.architectureData
        or {}
    )


    # --------------------------------------------------------
    # Components
    # --------------------------------------------------------

    v1_components = (
        v1_data.get(
            "components",
            []
        )
    )

    v2_components = (
        v2_data.get(
            "components",
            []
        )
    )


    v1_comp_ids = {

        component.get("id")
        or component.get("name"):
            component

        for component
        in v1_components
    }


    v2_comp_ids = {

        component.get("id")
        or component.get("name"):
            component

        for component
        in v2_components
    }


    added_components = [

        component

        for key, component
        in v2_comp_ids.items()

        if key not in v1_comp_ids
    ]


    removed_components = [

        component

        for key, component
        in v1_comp_ids.items()

        if key not in v2_comp_ids
    ]


    # --------------------------------------------------------
    # APIs
    # --------------------------------------------------------

    v1_apis = (
        v1_data.get(
            "apis",
            []
        )
    )

    v2_apis = (
        v2_data.get(
            "apis",
            []
        )
    )


    v1_api_paths = {

        api.get("path"):
            api

        for api
        in v1_apis
    }


    v2_api_paths = {

        api.get("path"):
            api

        for api
        in v2_apis
    }


    added_apis = [

        api

        for key, api
        in v2_api_paths.items()

        if key not in v1_api_paths
    ]


    removed_apis = [

        api

        for key, api
        in v1_api_paths.items()

        if key not in v2_api_paths
    ]


    # --------------------------------------------------------
    # Database tables
    # --------------------------------------------------------

    v1_tables = (
        v1_data.get(
            "database_schema",
            []
        )
    )

    v2_tables = (
        v2_data.get(
            "database_schema",
            []
        )
    )


    v1_table_names = {

        table.get("name"):
            table

        for table
        in v1_tables
    }


    v2_table_names = {

        table.get("name"):
            table

        for table
        in v2_tables
    }


    added_tables = [

        table

        for key, table
        in v2_table_names.items()

        if key not in v1_table_names
    ]


    removed_tables = [

        table

        for key, table
        in v1_table_names.items()

        if key not in v2_table_names
    ]


    return {

        "v1": v1,

        "v2": v2,

        "added_components":
            added_components,

        "removed_components":
            removed_components,

        "added_apis":
            added_apis,

        "removed_apis":
            removed_apis,

        "added_tables":
            added_tables,

        "removed_tables":
            removed_tables,

        "summary": (
            f"Added {len(added_components)} components, "
            f"removed {len(removed_components)} components, "
            f"added {len(added_apis)} APIs, "
            f"removed {len(removed_apis)} APIs, "
            f"added {len(added_tables)} tables, "
            f"removed {len(removed_tables)} tables"
        )
    }